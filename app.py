from __future__ import annotations

import hashlib
import html
import importlib.util
import json
import os
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone

MALAYSIA_TZ = timezone(timedelta(hours=8))
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
import plotly.express as px
import streamlit as st

from loader import Inbox


# -----------------------------------------------------------------------------
# Configuration
# -----------------------------------------------------------------------------

st.set_page_config(
    page_title="Shipping Document Verification",
    page_icon="🚢",
    layout="wide",
    initial_sidebar_state="collapsed",
)

PROJECT_ROOT = Path(__file__).resolve().parent
SRC_DIRECTORY = PROJECT_ROOT / "src"
OUTPUT_FILE = PROJECT_ROOT / "output" / "submission.json"
ROOT_OUTPUT_FILE = PROJECT_ROOT / "submission.json"

if str(SRC_DIRECTORY) not in sys.path:
    sys.path.insert(0, str(SRC_DIRECTORY))

try:
    from sdoc_pipeline import FIELDS as PIPELINE_FIELDS
    from sdoc_pipeline import _normalized as pipeline_normalized
    from sdoc_pipeline import extract_attachment, parse_document, process

    PIPELINE_AVAILABLE = True
    PIPELINE_IMPORT_ERROR = None
except Exception as import_error:  # pragma: no cover - deployment safeguard
    PIPELINE_FIELDS = ()
    pipeline_normalized = None
    extract_attachment = None
    parse_document = None
    process = None
    PIPELINE_AVAILABLE = False
    PIPELINE_IMPORT_ERROR = str(import_error)

VALID_CATEGORIES = {
    "BL_COMPARISON",
    "SI_REQUEST",
    "INVOICE_QUERY",
    "GENERAL",
    "SPAM",
}

VALID_STATUSES = {"OK", "MISMATCH", "NEEDS_REVIEW"}

VALID_REVIEW_REASONS = {
    "wrong_doc_type",
    "missing_attachment",
    "unreadable",
    "missing_value",
}

COMPARISON_FIELDS = [
    "shipper",
    "consignee",
    "notify_party",
    "port_of_loading",
    "port_of_discharge",
    "container_count",
    "gross_weight_kg",
]

MANUAL_UPLOAD_TYPES = ["json", "csv", "txt", "pdf", "docx", "xlsx"]

if PIPELINE_FIELDS and list(PIPELINE_FIELDS) != COMPARISON_FIELDS:
    raise RuntimeError("Dashboard and pipeline comparison fields do not match.")

CATEGORY_NAMES = {
    "BL_COMPARISON": "BL Comparison",
    "SI_REQUEST": "SI Request",
    "INVOICE_QUERY": "Invoice Query",
    "GENERAL": "General",
    "SPAM": "Spam",
    "NOT_PROCESSED": "Not Processed",
}

STATUS_NAMES = {
    "OK": "No Mismatch",
    "MISMATCH": "Mismatch",
    "NEEDS_REVIEW": "Needs Review",
    "NOT_PROCESSED": "Not Processed",
    "NOT_APPLICABLE": "Not Applicable",
}

STATUS_ICONS = {
    "OK": "✓",
    "MISMATCH": "⚠",
    "NEEDS_REVIEW": "!",
    "NOT_PROCESSED": "·",
    "NOT_APPLICABLE": "·",
}

CATEGORY_COLOURS = {
    "BL Comparison": "#F4A62A",
    "SI Request": "#6C89FF",
    "Invoice Query": "#A879F7",
    "General": "#7E8A9A",
    "Spam": "#F35D77",
    "Not Processed": "#4C5868",
}

STATUS_COLOURS = {
    "No Mismatch": "#33D69F",
    "Mismatch": "#F35D77",
    "Needs Review": "#F4A62A",
    "Not Processed": "#697689",
    "Not Applicable": "#697689",
}


# -----------------------------------------------------------------------------
# Styling - one consolidated theme to avoid conflicting CSS layers
# -----------------------------------------------------------------------------

st.markdown(
    """
    <style>
        :root {
            --orange: #F4A62A;
            --orange-dark: #D98A13;
            --page: #080D12;
            --panel: #111821;
            --panel-hover: #151E29;
            --border: #24303D;
            --text: #F4F7FB;
            --soft: #A6B0BF;
            --muted: #748195;
            --green: #33D69F;
            --red: #F35D77;
            --blue: #6C89FF;
        }

        html, body, [class*="css"] {
            font-family: Inter, ui-sans-serif, system-ui, -apple-system,
                BlinkMacSystemFont, "Segoe UI", sans-serif;
        }

        .stApp, [data-testid="stAppViewContainer"] {
            background: var(--page);
            color: var(--text);
        }

        [data-testid="stHeader"] {
            background: rgba(8, 13, 18, 0.94);
            border-bottom: 1px solid rgba(36, 48, 61, 0.7);
        }

        [data-testid="stSidebar"] { display: none; }

        .block-container {
            max-width: 1500px;
            padding: 1rem 1.25rem 4rem;
        }

        h1, h2, h3, h4, p, label, li { color: var(--text); }
        h2, h3 { letter-spacing: -0.25px; }
        a { color: var(--orange); }
        hr { border-color: var(--border); }
        code { color: #FFD18A; background: #161D26; }

        .brand-row {
            display: flex;
            align-items: center;
            justify-content: space-between;
            gap: 24px;
            margin: 0 0 10px;
        }

        .wordmark {
            color: #FFFFFF;
            font-size: 25px;
            font-weight: 900;
            font-style: italic;
            letter-spacing: -1.2px;
        }

        .wordmark-accent {
            display: inline-block;
            width: 19px;
            height: 5px;
            margin-right: -18px;
            margin-bottom: 19px;
            border-radius: 50% 50% 0 0;
            background: var(--orange);
            transform: rotate(-7deg);
        }

        .health-pill {
            display: inline-flex;
            align-items: center;
            gap: 8px;
            padding: 7px 12px;
            color: var(--soft);
            background: var(--panel);
            border: 1px solid var(--border);
            border-radius: 999px;
            font-size: 12px;
            font-weight: 700;
        }

        .health-dot {
            width: 8px;
            height: 8px;
            border-radius: 50%;
            background: var(--health-colour, var(--orange));
            box-shadow: 0 0 0 4px color-mix(in srgb, var(--health-colour, var(--orange)) 14%, transparent);
        }

        .st-key-navigation { margin: 0 0 12px; }

        .st-key-navigation div[role="radiogroup"] {
            display: flex;
            flex-flow: row wrap;
            gap: 8px;
        }

        .st-key-navigation div[role="radiogroup"] label {
            min-height: 37px;
            margin: 0;
            padding: 8px 16px;
            color: var(--soft);
            background: transparent;
            border: 1px solid var(--border);
            border-radius: 999px;
            cursor: pointer;
            transition: 160ms ease;
        }

        .st-key-navigation div[role="radiogroup"] label > div:first-child,
        .st-key-navigation [data-baseweb="radio"] > div:first-child {
            display: none !important;
        }

        .st-key-navigation div[role="radiogroup"] label p {
            margin: 0;
            color: inherit !important;
            font-size: 13px;
            font-weight: 650;
            white-space: nowrap;
        }

        .st-key-navigation div[role="radiogroup"] label:hover {
            color: var(--orange);
            background: rgba(244, 166, 42, 0.08);
            border-color: rgba(244, 166, 42, 0.5);
            transform: translateY(-1px);
        }

        .st-key-navigation div[role="radiogroup"] label:has(input:checked) {
            color: #111820;
            background: var(--orange);
            border-color: var(--orange);
            box-shadow: 0 5px 16px rgba(244, 166, 42, 0.2);
        }

        .top-divider {
            height: 1px;
            margin: 4px 0 28px;
            background: var(--border);
        }

        .page-title {
            margin: 0 0 5px;
            color: var(--text);
            font-size: clamp(31px, 3vw, 43px);
            font-weight: 850;
            letter-spacing: -1.2px;
            line-height: 1.08;
        }

        .page-subtitle {
            margin: 0 0 22px;
            color: var(--muted);
            font-size: 14px;
            line-height: 1.55;
        }

        .eyebrow {
            margin-bottom: 8px;
            color: var(--orange);
            font-size: 11px;
            font-weight: 800;
            letter-spacing: 1.2px;
            text-transform: uppercase;
        }

        [data-testid="stMetric"] {
            min-height: 112px;
            padding: 17px 17px 15px;
            background: var(--panel);
            border: 1px solid var(--border);
            border-left: 4px solid var(--orange);
            border-radius: 11px;
            box-shadow: 0 10px 28px rgba(0, 0, 0, 0.14);
            transition: 160ms ease;
        }

        [data-testid="stMetric"]:hover {
            background: var(--panel-hover);
            border-color: #344152;
            border-left-color: var(--orange);
            transform: translateY(-2px);
        }

        [data-testid="stMetricLabel"],
        [data-testid="stMetricLabel"] p {
            color: var(--soft) !important;
            font-size: 13px;
            font-weight: 600;
        }

        [data-testid="stMetricValue"],
        [data-testid="stMetricValue"] div {
            color: var(--text) !important;
            font-size: 29px;
            font-weight: 800;
        }

        div.stButton > button,
        div.stDownloadButton > button {
            min-height: 41px;
            padding: 8px 18px;
            color: #10161E;
            background: var(--orange);
            border: 1px solid var(--orange);
            border-radius: 9px;
            font-size: 13px;
            font-weight: 750;
            transition: 160ms ease;
        }

        div.stButton > button:hover,
        div.stDownloadButton > button:hover {
            color: #FFFFFF;
            background: var(--orange-dark);
            border-color: var(--orange-dark);
            transform: translateY(-1px);
            box-shadow: 0 7px 18px rgba(244, 166, 42, 0.18);
        }

        div.stButton > button:disabled {
            color: #667181;
            background: #1A222D;
            border-color: #252F3C;
        }

        [data-testid="stTextInput"] input,
        [data-testid="stTextArea"] textarea,
        [data-baseweb="select"] > div,
        [data-testid="stFileUploaderDropzone"] {
            color: var(--text) !important;
            background: var(--panel) !important;
            border-color: var(--border) !important;
            border-radius: 9px !important;
        }

        [data-testid="stTextArea"] textarea:disabled {
            color: #CFD6E1 !important;
            -webkit-text-fill-color: #CFD6E1 !important;
            opacity: 1;
        }

        [data-testid="stExpander"] {
            overflow: hidden;
            color: var(--text);
            background: var(--panel);
            border: 1px solid var(--border);
            border-radius: 11px;
        }

        [data-testid="stDataFrame"], [data-testid="stPlotlyChart"] {
            overflow: hidden;
            background: var(--panel);
            border: 1px solid var(--border);
            border-radius: 11px;
        }

        [data-testid="stPlotlyChart"] { padding: 8px; }

        [data-testid="stAlert"] {
            color: var(--text);
            border-radius: 9px;
        }

        [data-baseweb="tab-list"] {
            gap: 22px;
            border-bottom: 1px solid var(--border);
        }

        [data-baseweb="tab"] {
            color: var(--soft);
            background: transparent;
            font-size: 13px;
            font-weight: 650;
        }

        [data-baseweb="tab"]:hover,
        [aria-selected="true"][data-baseweb="tab"] { color: var(--orange); }
        [data-baseweb="tab-highlight"] { background: var(--orange); }

        .status-banner {
            padding: 14px 16px;
            color: var(--text);
            background: var(--panel);
            border: 1px solid var(--border);
            border-left: 4px solid var(--status-colour, var(--muted));
            border-radius: 9px;
            font-size: 13px;
            font-weight: 700;
        }

        .format-badge, .status-badge {
            display: inline-block;
            margin-left: 6px;
            padding: 3px 8px;
            color: var(--soft);
            background: #19222D;
            border: 1px solid var(--border);
            border-radius: 999px;
            font-size: 10px;
            font-weight: 800;
            letter-spacing: 0.4px;
        }

        .source-evidence {
            margin-top: 8px;
            padding: 10px 12px;
            color: var(--soft);
            background: #0C1219;
            border-left: 3px solid #405064;
            border-radius: 6px;
            font-size: 12px;
            line-height: 1.5;
        }

        .review-card {
            padding: 15px 16px;
            background: var(--panel);
            border: 1px solid var(--border);
            border-radius: 11px;
        }

        @media (max-width: 900px) {
            .block-container { padding: 0.8rem 0.8rem 3rem; }
            .health-pill { display: none; }
            .st-key-navigation div[role="radiogroup"] label { padding: 7px 10px; }
            .st-key-navigation div[role="radiogroup"] label p { font-size: 12px; }
            [data-testid="stMetric"] { min-height: 98px; }
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# -----------------------------------------------------------------------------
# State and data helpers
# -----------------------------------------------------------------------------

SESSION_DEFAULTS = {
    "submission": {},
    "submission_source": "No results",
    "submission_generated_at": None,
    "submission_source_key": None,
    "review_history": [],
    "review_flash": None,
    "manual_comparison": None,
    "comparison_email": None,
    "review_email": None,
    "navigation": "Dashboard",
}

for state_key, default_value in SESSION_DEFAULTS.items():
    if state_key not in st.session_state:
        st.session_state[state_key] = default_value


@st.cache_resource
def create_inbox(source: str) -> Inbox:
    return Inbox(source)


@st.cache_data(show_spinner=False)
def load_emails(source: str) -> list[dict[str, Any]]:
    return Inbox(source).emails()


@st.cache_data(show_spinner=False)
def load_json_file(path_text: str, modified_time: float) -> dict[str, dict[str, Any]]:
    del modified_time
    path = Path(path_text)
    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)
    if not isinstance(data, dict):
        raise ValueError("The result must be one JSON object keyed by email ID.")
    return data


def uploaded_json(uploaded_file: Any) -> tuple[dict[str, dict[str, Any]], str]:
    raw = uploaded_file.getvalue()
    data = json.loads(raw.decode("utf-8"))
    if not isinstance(data, dict):
        raise ValueError("The uploaded JSON must be keyed by email ID.")
    return data, hashlib.sha256(raw).hexdigest()


def load_default_submission(source_key: str) -> None:
    if st.session_state.submission_source_key == source_key:
        return

    for candidate in (OUTPUT_FILE, ROOT_OUTPUT_FILE):
        if candidate.exists():
            st.session_state.submission = load_json_file(
                str(candidate), candidate.stat().st_mtime
            )
            st.session_state.submission_source = str(candidate.relative_to(PROJECT_ROOT))
            st.session_state.submission_generated_at = datetime.fromtimestamp(
                candidate.stat().st_mtime, tz=timezone.utc
            ).isoformat()
            st.session_state.submission_source_key = source_key
            return

    st.session_state.submission = {}
    st.session_state.submission_source = "No results"
    st.session_state.submission_generated_at = None
    st.session_state.submission_source_key = source_key


def normalise_result(email_id: str, result: Any) -> dict[str, Any]:
    if not isinstance(result, dict):
        return {
            "email_id": email_id,
            "category": "NOT_PROCESSED",
            "status": "NOT_PROCESSED",
            "review_reason": None,
            "defect_fields": [],
            "has_defect": False,
        }

    category = result.get("category", "NOT_PROCESSED")
    status = result.get("status") or (
        "NOT_PROCESSED" if category == "BL_COMPARISON" else "NOT_APPLICABLE"
    )
    fields = result.get("defect_fields") or []
    if not isinstance(fields, list):
        fields = []

    return {
        **result,
        "email_id": email_id,
        "category": category,
        "status": status,
        "review_reason": result.get("review_reason"),
        "defect_fields": fields,
        "has_defect": bool(result.get("has_defect", False)),
    }


def results_dataframe(
    emails: list[dict[str, Any]], submission: dict[str, dict[str, Any]]
) -> pd.DataFrame:
    rows = []
    for email in emails:
        email_id = email["email_id"]
        result = normalise_result(email_id, submission.get(email_id))
        rows.append(
            {
                "Email ID": email_id,
                "Sender": email.get("from", ""),
                "Subject": email.get("subject", ""),
                "Attachments": len(email.get("attachments", [])),
                "Category": result["category"],
                "Category Display": CATEGORY_NAMES.get(
                    result["category"], result["category"].replace("_", " ").title()
                ),
                "Status": result["status"],
                "Status Display": STATUS_NAMES.get(
                    result["status"], result["status"].replace("_", " ").title()
                ),
                "Has Defect": result["has_defect"],
                "Defect Fields": ", ".join(result["defect_fields"]),
                "Issue Count": len(result["defect_fields"]),
                "Review Reason": result["review_reason"] or "",
            }
        )
    return pd.DataFrame(rows)


def validate_submission(
    emails: list[dict[str, Any]], submission: dict[str, dict[str, Any]]
) -> list[str]:
    errors: list[str] = []
    expected_ids = {email["email_id"] for email in emails}
    actual_ids = set(submission)

    missing_ids = expected_ids - actual_ids
    extra_ids = actual_ids - expected_ids
    if missing_ids:
        errors.append(f"{len(missing_ids)} email IDs are missing.")
    if extra_ids:
        errors.append(f"{len(extra_ids)} unexpected email IDs are present.")

    for email_id in sorted(expected_ids & actual_ids):
        result = submission[email_id]
        if not isinstance(result, dict):
            errors.append(f"{email_id}: result must be a JSON object.")
            continue

        category = result.get("category")
        status = result.get("status")
        has_defect = result.get("has_defect")
        defect_fields = result.get("defect_fields")
        review_reason = result.get("review_reason")

        if category not in VALID_CATEGORIES:
            errors.append(f"{email_id}: invalid category {category!r}.")
        if status not in VALID_STATUSES:
            errors.append(f"{email_id}: invalid status {status!r}.")
        if not isinstance(has_defect, bool):
            errors.append(f"{email_id}: has_defect must be true or false.")
        if not isinstance(defect_fields, list):
            errors.append(f"{email_id}: defect_fields must be a list.")
            defect_fields = []

        invalid_fields = [field for field in defect_fields if field not in COMPARISON_FIELDS]
        if invalid_fields:
            errors.append(f"{email_id}: invalid defect fields {invalid_fields}.")
        if status == "MISMATCH" and (not has_defect or not defect_fields):
            errors.append(f"{email_id}: MISMATCH requires defect fields.")
        if status == "OK" and (has_defect or defect_fields):
            errors.append(f"{email_id}: OK cannot contain defects.")
        if status == "NEEDS_REVIEW" and review_reason not in VALID_REVIEW_REASONS:
            errors.append(f"{email_id}: NEEDS_REVIEW requires a valid review reason.")
        if status != "NEEDS_REVIEW" and review_reason is not None:
            errors.append(f"{email_id}: review_reason must be null for {status}.")

    return errors


def attachment_roles(email: dict[str, Any]) -> tuple[list[str], list[str]]:
    paths = [str(path) for path in email.get("attachments", [])]
    si_paths = [path for path in paths if "_SI" in Path(path).stem.upper()]
    bl_paths = [path for path in paths if "_BL" in Path(path).stem.upper()]
    return si_paths, bl_paths


@st.cache_data(show_spinner=False)
def comparison_evidence(source: str, email_id: str) -> dict[str, Any]:
    if not PIPELINE_AVAILABLE or extract_attachment is None:
        return {"error": PIPELINE_IMPORT_ERROR or "Pipeline unavailable"}

    local_inbox = Inbox(source)
    email = next((item for item in local_inbox.emails() if item["email_id"] == email_id), None)
    if email is None:
        return {"error": "Email not found"}

    si_paths, bl_paths = attachment_roles(email)
    if not si_paths or not bl_paths:
        return {"error": "Required SI or BL attachment is missing"}

    si_result = extract_attachment(si_paths[0], local_inbox.read_bytes(si_paths[0]))
    bl_result = extract_attachment(bl_paths[0], local_inbox.read_bytes(bl_paths[0]))

    if si_result.document is None or bl_result.document is None:
        return {
            "error": si_result.error or bl_result.error or "Document could not be read",
            "si_format": si_result.format,
            "bl_format": bl_result.format,
        }

    rows = []
    for field in COMPARISON_FIELDS:
        si_value = si_result.document.values.get(field, "")
        bl_value = bl_result.document.values.get(field, "")
        if not si_value or not bl_value:
            outcome = "Evidence unavailable"
        else:
            si_normalized = pipeline_normalized(field, si_value)
            bl_normalized = pipeline_normalized(field, bl_value)
            outcome = "Match" if si_normalized == bl_normalized else "Mismatch"

        rows.append(
            {
                "Field Key": field,
                "Field": field.replace("_", " ").title(),
                "SI Value": si_value or "Not extracted",
                "BL Value": bl_value or "Not extracted",
                "Outcome": outcome,
                "SI Evidence": si_result.document.evidence.get(field, "Not available"),
                "BL Evidence": bl_result.document.evidence.get(field, "Not available"),
            }
        )

    return {
        "error": None,
        "rows": rows,
        "si_format": si_result.format.upper(),
        "bl_format": bl_result.format.upper(),
        "si_path": si_paths[0],
        "bl_path": bl_paths[0],
    }


@st.cache_data(show_spinner=False)
def attachment_preview_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".csv"}:
        return content.decode("utf-8", errors="replace")

    if suffix == ".json":
        try:
            return json.dumps(json.loads(content.decode("utf-8")), indent=2, ensure_ascii=False)
        except (UnicodeDecodeError, json.JSONDecodeError):
            return content.decode("utf-8", errors="replace")

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            return "\n\n".join(
                page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages
            ).strip()
        except Exception:
            return ""

    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(parts)
        except Exception:
            return ""

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            rows = []
            for worksheet in workbook.worksheets:
                rows.append(f"[{worksheet.title}]")
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        rows.append(" | ".join(values))
            return "\n".join(rows)
        except Exception:
            return ""

    return ""


def csv_document_text(content: bytes) -> str:
    """Convert common field/value or one-record CSV layouts into label text."""
    try:
        dataframe = pd.read_csv(BytesIO(content), dtype=str, keep_default_na=False)
    except UnicodeDecodeError:
        dataframe = pd.read_csv(
            BytesIO(content), dtype=str, keep_default_na=False, encoding="latin-1"
        )

    if dataframe.empty and not len(dataframe.columns):
        return ""

    known_header_tokens = {
        "shipper",
        "consignee",
        "notify_party",
        "port_of_loading",
        "port_of_discharge",
        "container_count",
        "gross_weight_kg",
        "notify party",
        "port of loading",
        "port of discharge",
        "container count",
        "gross weight kg",
    }
    normalized_headers = {
        str(column).strip().lower().replace("-", "_") for column in dataframe.columns
    }

    # Layout 1: the seven fields are columns and the first record holds values.
    if normalized_headers & known_header_tokens and not dataframe.empty:
        first_record = dataframe.iloc[0]
        return "\n".join(
            f"{column}: {first_record[column]}" for column in dataframe.columns
        )

    # Layout 2: every row is label,value (additional value columns are joined).
    rows: list[str] = []
    for _, record in dataframe.iterrows():
        values = [str(value).strip() for value in record.tolist() if str(value).strip()]
        if len(values) >= 2:
            rows.append(f"{values[0]}: {' '.join(values[1:])}")

    # pandas treats the first CSV row as headers, so preserve it for headerless
    # field/value files as well.
    columns = [str(column).strip() for column in dataframe.columns]
    if len(columns) >= 2 and not normalized_headers & known_header_tokens:
        rows.insert(0, f"{columns[0]}: {' '.join(columns[1:])}")

    return "\n".join(rows)


def json_document_text(content: bytes) -> str:
    """Convert a JSON shipment object into the pipeline's label/value text."""
    data = json.loads(content.decode("utf-8"))

    if isinstance(data, list):
        if not data or not isinstance(data[0], dict):
            return ""
        data = data[0]

    if not isinstance(data, dict):
        return ""

    # Accept wrappers commonly used for SI/BL documents.
    for wrapper in ("shipping_instruction", "si", "bill_of_lading", "bl", "document", "data"):
        candidate = data.get(wrapper)
        if isinstance(candidate, dict):
            data = candidate
            break

    return "\n".join(
        f"{str(field).replace('_', ' ')}: {value}"
        for field, value in data.items()
        if value is not None and not isinstance(value, (dict, list))
    )


def extract_uploaded_document(
    filename: str, content: bytes
) -> tuple[Any | None, str, str | None]:
    """Extract one manually uploaded document with CSV support added locally."""
    if not PIPELINE_AVAILABLE:
        return None, Path(filename).suffix.lstrip(".").upper(), PIPELINE_IMPORT_ERROR

    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        try:
            text = json_document_text(content)
            if not text.strip():
                return None, "JSON", "The JSON contains no readable shipment fields."
            return parse_document(text), "JSON", None
        except Exception as error:
            return None, "JSON", f"JSON extraction failed: {error}"

    if suffix == ".csv":
        try:
            text = csv_document_text(content)
            if not text.strip():
                return None, "CSV", "The CSV contains no readable field values."
            return parse_document(text), "CSV", None
        except Exception as error:
            return None, "CSV", f"CSV extraction failed: {error}"

    result = extract_attachment(filename, content)
    return result.document, result.format.upper(), result.error


def compare_uploaded_documents(
    si_name: str,
    si_content: bytes,
    bl_name: str,
    bl_content: bytes,
) -> dict[str, Any]:
    si_document, si_format, si_error = extract_uploaded_document(si_name, si_content)
    bl_document, bl_format, bl_error = extract_uploaded_document(bl_name, bl_content)

    if si_document is None or bl_document is None:
        return {
            "status": "NEEDS_REVIEW",
            "review_reason": "unreadable",
            "defect_fields": [],
            "has_defect": False,
            "error": si_error or bl_error or "A document could not be read.",
            "si_format": si_format,
            "bl_format": bl_format,
            "rows": [],
        }

    rows = []
    missing_fields = []
    defect_fields = []
    for field in COMPARISON_FIELDS:
        si_value = si_document.values.get(field, "")
        bl_value = bl_document.values.get(field, "")
        if not si_value or not bl_value:
            outcome = "Evidence unavailable"
            missing_fields.append(field)
        else:
            match = pipeline_normalized(field, si_value) == pipeline_normalized(field, bl_value)
            outcome = "Match" if match else "Mismatch"
            if not match:
                defect_fields.append(field)

        rows.append(
            {
                "Field Key": field,
                "Field": field.replace("_", " ").title(),
                "SI Value": si_value or "Not extracted",
                "BL Value": bl_value or "Not extracted",
                "Outcome": outcome,
                "SI Evidence": si_document.evidence.get(field, "Not available"),
                "BL Evidence": bl_document.evidence.get(field, "Not available"),
            }
        )

    if missing_fields:
        status = "NEEDS_REVIEW"
        review_reason = "missing_value"
        final_defects = []
    elif defect_fields:
        status = "MISMATCH"
        review_reason = None
        final_defects = defect_fields
    else:
        status = "OK"
        review_reason = None
        final_defects = []

    return {
        "status": status,
        "review_reason": review_reason,
        "defect_fields": final_defects,
        "has_defect": status == "MISMATCH",
        "error": None,
        "si_format": si_format,
        "bl_format": bl_format,
        "missing_fields": missing_fields,
        "rows": rows,
    }


def safe_mime(extension: str) -> str:
    return {
        ".json": "application/json",
        ".csv": "text/csv",
        ".txt": "text/plain",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }.get(extension, "application/octet-stream")


@st.cache_data(show_spinner=False)
def render_pdf_pages(content: bytes, maximum_pages: int = 30) -> tuple[list[bytes], int, str | None]:
    """Render PDF pages to PNG so browser PDF security settings cannot block them."""
    try:
        import fitz

        document = fitz.open(stream=content, filetype="pdf")
        total_pages = len(document)
        matrix = fitz.Matrix(1.55, 1.55)
        pages = [
            document.load_page(index).get_pixmap(matrix=matrix, alpha=False).tobytes("png")
            for index in range(min(total_pages, maximum_pages))
        ]
        document.close()
        return pages, total_pages, None
    except Exception as error:
        return [], 0, str(error)


def display_pdf_preview(content: bytes, key_prefix: str) -> None:
    """Use the native viewer when installed; otherwise show rendered page images."""
    native_viewer_available = (
        hasattr(st, "pdf") and importlib.util.find_spec("streamlit_pdf") is not None
    )
    if native_viewer_available:
        st.pdf(content, height=620, key=f"pdf_{key_prefix}")
        return

    pages, total_pages, error = render_pdf_pages(content)
    if not pages:
        st.warning(
            "A visual PDF preview could not be rendered. "
            "The extracted text preview and download remain available."
        )
        if error:
            st.caption(f"Preview detail: {error}")
        return

    selected_page = 1
    if len(pages) > 1:
        selected_page = st.selectbox(
            "Preview page",
            options=list(range(1, len(pages) + 1)),
            key=f"pdf_page_{key_prefix}",
        )
    st.image(
        pages[selected_page - 1],
        caption=f"Page {selected_page} of {total_pages}",
        width="stretch",
    )
    if total_pages > len(pages):
        st.caption(
            f"The preview is limited to the first {len(pages)} pages. "
            "Download the PDF to view the remaining pages."
        )


# -----------------------------------------------------------------------------
# Presentation helpers
# -----------------------------------------------------------------------------

def page_header(title: str, subtitle: str, eyebrow: str | None = None) -> None:
    if eyebrow:
        st.markdown(f'<div class="eyebrow">{eyebrow}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-title">{title}</div>', unsafe_allow_html=True)
    st.markdown(f'<div class="page-subtitle">{subtitle}</div>', unsafe_allow_html=True)


def status_banner(result: dict[str, Any]) -> None:
    status = result["status"]
    colour = {
        "OK": "#33D69F",
        "MISMATCH": "#F35D77",
        "NEEDS_REVIEW": "#F4A62A",
    }.get(status, "#748195")

    if status == "OK":
        message = "No mismatch detected. All seven required fields match."
    elif status == "MISMATCH":
        message = f"Mismatch detected. {len(result.get('defect_fields', []))} field(s) require attention."
    elif status == "NEEDS_REVIEW":
        reason = str(result.get("review_reason") or "uncertain result").replace("_", " ")
        message = f"Human review required: {reason}."
    else:
        message = "This email has not been processed."

    st.markdown(
        f'<div class="status-banner" style="--status-colour:{colour}">'
        f'{STATUS_ICONS.get(status, "·")} {message}</div>',
        unsafe_allow_html=True,
    )


def format_plot(figure: Any, *, height: int = 360) -> Any:
    figure.update_layout(
        height=height,
        margin=dict(l=20, r=20, t=20, b=20),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font_color="#A6B0BF",
        xaxis=dict(gridcolor="#24303D"),
        yaxis=dict(gridcolor="#24303D"),
    )
    return figure


def display_attachment(inbox: Inbox, path: str, key_prefix: str) -> None:
    filename = Path(path).name
    extension = Path(path).suffix.lower()
    try:
        content = inbox.read_bytes(path)
    except Exception as error:
        st.error(f"Could not open {filename}: {error}")
        return

    heading_col, download_col = st.columns([4, 1], vertical_alignment="center")
    with heading_col:
        st.markdown(
            f"**{filename}** <span class='format-badge'>{extension.lstrip('.').upper()}</span>",
            unsafe_allow_html=True,
        )
    with download_col:
        st.download_button(
            "Download",
            data=content,
            file_name=filename,
            mime=safe_mime(extension),
            key=f"download_{key_prefix}_{filename}",
            width="stretch",
        )

    if extension == ".pdf":
        display_pdf_preview(content, f"stored_{key_prefix}_{filename}")

    preview = attachment_preview_text(filename, content)
    if preview:
        label = "Extracted text" if extension != ".txt" else "Document contents"
        with st.expander(label, expanded=extension != ".pdf"):
            st.text_area(
                label,
                value=preview,
                height=330,
                disabled=True,
                label_visibility="collapsed",
                key=f"preview_{key_prefix}_{filename}",
            )
    elif extension != ".pdf":
        st.warning("A safe inline preview could not be generated. Download the original document to inspect it.")


def display_uploaded_attachment(uploaded_file: Any, key_prefix: str) -> None:
    """Preview and download a document selected through a manual uploader."""
    filename = uploaded_file.name
    content = uploaded_file.getvalue()
    extension = Path(filename).suffix.lower()

    heading_col, download_col = st.columns([4, 1], vertical_alignment="center")
    with heading_col:
        st.markdown(
            f"**{html.escape(filename)}** "
            f"<span class='format-badge'>{extension.lstrip('.').upper()}</span>",
            unsafe_allow_html=True,
        )
    with download_col:
        st.download_button(
            "Download",
            data=content,
            file_name=filename,
            mime=safe_mime(extension),
            key=f"manual_download_{key_prefix}_{filename}",
            width="stretch",
        )

    if extension == ".pdf":
        display_pdf_preview(content, f"uploaded_{key_prefix}_{filename}")

    preview = attachment_preview_text(filename, content)
    if preview:
        with st.expander("Document preview", expanded=extension != ".pdf"):
            st.text_area(
                "Document preview",
                value=preview,
                height=300,
                disabled=True,
                label_visibility="collapsed",
                key=f"manual_preview_{key_prefix}_{filename}",
            )
    elif extension != ".pdf":
        st.warning("The file was uploaded, but a readable preview could not be generated.")


def comparison_table_data(
    source: str, email_id: str, result: dict[str, Any]
) -> tuple[pd.DataFrame, dict[str, Any]]:
    evidence = comparison_evidence(source, email_id)
    if not evidence.get("error"):
        rows = evidence["rows"]
    else:
        defect_fields = set(result.get("defect_fields", []))
        rows = []
        for field in COMPARISON_FIELDS:
            rows.append(
                {
                    "Field Key": field,
                    "Field": field.replace("_", " ").title(),
                    "SI Value": "Evidence unavailable",
                    "BL Value": "Evidence unavailable",
                    "Outcome": "Mismatch" if field in defect_fields else "Evidence unavailable",
                    "SI Evidence": "Not available",
                    "BL Evidence": "Not available",
                }
            )
    return pd.DataFrame(rows), evidence


def evidence_details(table: pd.DataFrame) -> None:
    mismatches = table[table["Outcome"].isin(["Mismatch", "Evidence unavailable"])]
    if mismatches.empty:
        st.caption("All extracted values were available and matched after normalization.")
        return

    with st.expander("Source evidence and comparison explanation", expanded=True):
        for _, row in mismatches.iterrows():
            st.markdown(f"**{row['Field']} — {row['Outcome']}**")
            st.markdown(
                f'<div class="source-evidence"><b>SI:</b> {html.escape(str(row["SI Evidence"]))}<br>'
                f'<b>BL:</b> {html.escape(str(row["BL Evidence"]))}</div>',
                unsafe_allow_html=True,
            )
        st.caption(
            "Names and ports are compared after capitalization, spacing, and punctuation normalization. "
            "Counts and weights are compared as normalized numeric values."
        )


def submission_download(submission: dict[str, dict[str, Any]], key: str) -> None:
    st.download_button(
        "Download Results JSON",
        data=json.dumps(submission, indent=2, ensure_ascii=False) + "\n",
        file_name="submission.json",
        mime="application/json",
        key=key,
        width="stretch",
    )


def navigate(page_name: str, email_id: str | None = None) -> None:
    st.session_state.navigation = page_name
    if email_id:
        if page_name == "Compare Documents":
            st.session_state.comparison_email = email_id
        elif page_name == "Review Queue":
            st.session_state.review_email = email_id


# -----------------------------------------------------------------------------
# Header, controls, and loading
# -----------------------------------------------------------------------------

default_source = os.getenv("DATA_SOURCE", str(PROJECT_ROOT))
allow_custom_source = os.getenv("ALLOW_CUSTOM_DATA_SOURCE", "false").lower() == "true"

with st.expander("Data and developer controls", expanded=False):
    control_columns = st.columns([2.4, 1.7, 0.8])
    with control_columns[0]:
        if allow_custom_source:
            data_source = st.text_input(
                "Data source",
                value=default_source,
                help="Use the bundle directory or an approved Inbox HTTP endpoint.",
            )
        else:
            data_source = default_source
            st.text_input(
                "Data source",
                value=data_source,
                disabled=True,
                help="Set ALLOW_CUSTOM_DATA_SOURCE=true to edit this value.",
            )
    with control_columns[1]:
        uploaded_submission = st.file_uploader(
            "Load results JSON",
            type=["json"],
            help="Optional. Uploaded results stay in this browser session.",
        )
    with control_columns[2]:
        refresh_requested = st.button("Refresh data", width="stretch")

if refresh_requested:
    st.cache_data.clear()
    st.cache_resource.clear()
    st.session_state.submission_source_key = None
    st.rerun()

try:
    inbox = create_inbox(data_source)
    emails = load_emails(data_source)
except Exception as error:
    page_header("Shipping Document Verification", "The inbox could not be loaded.")
    st.error("Check that loader.py, inbox/, and attachments/ are available in the configured project directory.")
    with st.expander("Technical details"):
        st.exception(error)
    st.stop()

source_key = hashlib.sha256(data_source.encode("utf-8")).hexdigest()
load_default_submission(source_key)

if uploaded_submission is not None:
    try:
        upload_data, upload_hash = uploaded_json(uploaded_submission)
        upload_key = f"upload:{upload_hash}"
        if st.session_state.submission_source_key != upload_key:
            st.session_state.submission = upload_data
            st.session_state.submission_source = "Uploaded JSON"
            st.session_state.submission_generated_at = datetime.now(timezone.utc).isoformat()
            st.session_state.submission_source_key = upload_key
    except Exception as error:
        st.error(f"The uploaded result file is invalid: {error}")

submission: dict[str, dict[str, Any]] = st.session_state.submission
results_df = results_dataframe(emails, submission)
email_lookup = {email["email_id"]: email for email in emails}
validation_errors = validate_submission(emails, submission) if submission else []

if not submission:
    health_text = "Awaiting Results"
    health_colour = "#748195"
elif validation_errors:
    health_text = "Attention Required"
    health_colour = "#F4A62A"
else:
    health_text = "System Ready"
    health_colour = "#33D69F"

st.markdown(
    '<div class="brand-row">'
    '<div class="wordmark"><span class="wordmark-accent"></span>averis</div>'
    f'<div class="health-pill" style="--health-colour:{health_colour}">'
    f'<span class="health-dot"></span>{health_text}</div></div>',
    unsafe_allow_html=True,
)

page = st.radio(
    "Navigation",
    [
        "Dashboard",
        "Inbox",
        "Manual Upload",
        "Compare Documents",
        "Review Queue",
        "Results & Health",
    ],
    horizontal=True,
    label_visibility="collapsed",
    key="navigation",
)
st.markdown('<div class="top-divider"></div>', unsafe_allow_html=True)


# -----------------------------------------------------------------------------
# Dashboard
# -----------------------------------------------------------------------------

if page == "Dashboard":
    page_header(
        "Shipping Document Verification",
        "From inbox classification to evidence-backed discrepancy reporting.",
        "Operations overview",
    )

    action_col, meta_col = st.columns([1.2, 3.8], vertical_alignment="center")
    with action_col:
        if st.button(
            "▶ Run Verification",
            type="primary",
            width="stretch",
            disabled=not PIPELINE_AVAILABLE,
        ):
            try:
                progress = st.progress(10, text="Reading inbox and attachments...")
                with st.spinner("Classifying emails and comparing documents..."):
                    generated = process(inbox)
                progress.progress(100, text="Verification complete")
                st.session_state.submission = generated
                st.session_state.submission_source = "Generated in this session"
                st.session_state.submission_generated_at = datetime.now(timezone.utc).isoformat()
                st.session_state.submission_source_key = source_key
                st.rerun()
            except Exception as error:
                st.error(f"Verification could not be completed: {error}")

    with meta_col:
        generated_at = st.session_state.submission_generated_at
        timestamp_text = "Not generated"
        if generated_at:
            try:
                timestamp_text = (
                    datetime.fromisoformat(generated_at)
                    .astimezone(MALAYSIA_TZ)
                    .strftime("%d %b %Y, %H:%M MYT")
                )
            except ValueError:
                timestamp_text = generated_at
        st.caption(
            f"Result source: {st.session_state.submission_source} · Generated: {timestamp_text}"
        )

    if not PIPELINE_AVAILABLE:
        st.error(f"The processing pipeline could not be imported: {PIPELINE_IMPORT_ERROR}")

    processed = results_df[results_df["Category"] != "NOT_PROCESSED"]
    comparison_count = int((processed["Category"] == "BL_COMPARISON").sum())
    mismatch_count = int((processed["Status"] == "MISMATCH").sum())
    review_count = int((processed["Status"] == "NEEDS_REVIEW").sum())
    ok_comparisons = int(
        ((processed["Category"] == "BL_COMPARISON") & (processed["Status"] == "OK")).sum()
    )
    coverage = (len(processed) / len(emails) * 100) if emails else 0
    automation_rate = (
        (ok_comparisons + mismatch_count) / comparison_count * 100 if comparison_count else 0
    )

    metric_columns = st.columns(6)
    metric_columns[0].metric("Inbox", len(emails), help="Total email records loaded")
    metric_columns[1].metric("Coverage", f"{coverage:.1f}%", help="Emails with generated results")
    metric_columns[2].metric("BL Checks", comparison_count)
    metric_columns[3].metric("Mismatches", mismatch_count)
    metric_columns[4].metric("Needs Review", review_count)
    metric_columns[5].metric(
        "Automation", f"{automation_rate:.1f}%", help="BL checks completed without human review"
    )

    if not submission:
        st.warning("No results are loaded. Select Run Verification to process the inbox end to end.")
    else:
        st.write("")
        chart_left, chart_right = st.columns(2)
        with chart_left:
            st.subheader("Email Classification")
            category_counts = (
                processed["Category Display"].value_counts().rename_axis("Category").reset_index(name="Count")
            )
            figure = px.bar(
                category_counts,
                x="Category",
                y="Count",
                text="Count",
                color="Category",
                color_discrete_map=CATEGORY_COLOURS,
            )
            figure.update_layout(showlegend=False, xaxis_title="", yaxis_title="Emails")
            st.plotly_chart(format_plot(figure), width="stretch")

        with chart_right:
            st.subheader("Verification Outcomes")
            comparison_results = processed[processed["Category"] == "BL_COMPARISON"]
            status_counts = (
                comparison_results["Status Display"].value_counts().rename_axis("Status").reset_index(name="Count")
            )
            figure = px.bar(
                status_counts,
                x="Count",
                y="Status",
                orientation="h",
                text="Count",
                color="Status",
                color_discrete_map=STATUS_COLOURS,
            )
            figure.update_layout(showlegend=False, xaxis_title="Cases", yaxis_title="")
            st.plotly_chart(format_plot(figure), width="stretch")

        st.subheader("Demo Highlights")
        st.caption("Open one representative case to understand the system in under three minutes.")
        mismatch_ids = results_df.loc[results_df["Status"] == "MISMATCH", "Email ID"].tolist()
        ok_ids = results_df.loc[
            (results_df["Category"] == "BL_COMPARISON") & (results_df["Status"] == "OK"), "Email ID"
        ].tolist()
        review_ids = results_df.loc[results_df["Status"] == "NEEDS_REVIEW", "Email ID"].tolist()
        demo_columns = st.columns(3)
        demo_columns[0].button(
            "Inspect a mismatch",
            width="stretch",
            disabled=not mismatch_ids,
            on_click=navigate,
            args=("Compare Documents", mismatch_ids[0] if mismatch_ids else None),
        )
        demo_columns[1].button(
            "Inspect a clean match",
            width="stretch",
            disabled=not ok_ids,
            on_click=navigate,
            args=("Compare Documents", ok_ids[0] if ok_ids else None),
        )
        demo_columns[2].button(
            "Resolve a review case",
            width="stretch",
            disabled=not review_ids,
            on_click=navigate,
            args=("Review Queue", review_ids[0] if review_ids else None),
        )

        attention = results_df[results_df["Status"].isin(["MISMATCH", "NEEDS_REVIEW"])]
        st.subheader("Cases Requiring Attention")
        st.dataframe(
            attention[["Email ID", "Subject", "Status Display", "Issue Count", "Review Reason"]],
            width="stretch",
            hide_index=True,
            height=320,
            column_config={"Status Display": "Status", "Issue Count": "Issues"},
        )


# -----------------------------------------------------------------------------
# Inbox
# -----------------------------------------------------------------------------

elif page == "Inbox":
    page_header("Inbox", "Search every message and inspect its classification.", "Classification")

    filters = st.columns([2, 1, 1, 1])
    search = filters[0].text_input("Search", placeholder="Email ID, subject, or sender")
    category = filters[1].selectbox(
        "Category", ["All"] + sorted(results_df["Category Display"].unique().tolist())
    )
    status = filters[2].selectbox(
        "Status", ["All"] + sorted(results_df["Status Display"].unique().tolist())
    )
    attachment_filter = filters[3].selectbox(
        "Attachments", ["All", "Has attachments", "No attachments"]
    )

    filtered = results_df.copy()
    if search:
        mask = (
            filtered["Email ID"].str.contains(search, case=False, na=False)
            | filtered["Subject"].str.contains(search, case=False, na=False)
            | filtered["Sender"].str.contains(search, case=False, na=False)
        )
        filtered = filtered[mask]
    if category != "All":
        filtered = filtered[filtered["Category Display"] == category]
    if status != "All":
        filtered = filtered[filtered["Status Display"] == status]
    if attachment_filter == "Has attachments":
        filtered = filtered[filtered["Attachments"] > 0]
    elif attachment_filter == "No attachments":
        filtered = filtered[filtered["Attachments"] == 0]

    st.caption(f"{len(filtered)} email(s) displayed")
    st.dataframe(
        filtered[["Email ID", "Sender", "Subject", "Attachments", "Category Display", "Status Display"]],
        width="stretch",
        hide_index=True,
        height=430,
        column_config={"Category Display": "Category", "Status Display": "Status"},
    )

    if filtered.empty:
        st.info("No emails match the selected filters.")
    else:
        selected_id = st.selectbox("Open email", filtered["Email ID"].tolist())
        selected_email = email_lookup[selected_id]
        selected_result = normalise_result(selected_id, submission.get(selected_id))
        st.divider()
        st.subheader(selected_email.get("subject") or "No subject")
        details = st.columns(3)
        details[0].markdown(f"**Email ID:** {selected_id}")
        details[1].markdown(f"**Sender:** {selected_email.get('from', 'Unknown')}")
        details[2].markdown(f"**Category:** {CATEGORY_NAMES.get(selected_result['category'], selected_result['category'])}")

        email_tab, attachment_tab = st.tabs(["Email Body", "Attachments"])
        with email_tab:
            st.text_area("Message", selected_email.get("body", ""), height=350, disabled=True)
        with attachment_tab:
            attachments = selected_email.get("attachments", [])
            if not attachments:
                st.info("This email has no attachments.")
            for index, path in enumerate(attachments):
                display_attachment(inbox, path, f"inbox_{selected_id}_{index}")
                st.divider()


# -----------------------------------------------------------------------------
# Manual document upload and comparison
# -----------------------------------------------------------------------------

elif page == "Manual Upload":
    page_header(
        "Manual Document Check",
        "Upload a Shipping Instruction and draft Bill of Lading for an immediate seven-field comparison.",
        "Test your own documents",
    )

    st.info(
        "Supported formats: JSON, CSV, TXT, PDF, DOCX (Microsoft Word), and XLSX. "
        "Upload one SI and one draft BL, then select Compare Uploaded Documents."
    )

    upload_left, upload_right = st.columns(2)
    with upload_left:
        st.subheader("1. Shipping Instruction")
        manual_si = st.file_uploader(
            "Upload SI",
            type=MANUAL_UPLOAD_TYPES,
            key="manual_si_uploader",
            help="Accepted: JSON, CSV, TXT, PDF, DOCX, XLSX",
        )
        if manual_si is not None:
            display_uploaded_attachment(manual_si, "si")

    with upload_right:
        st.subheader("2. Draft Bill of Lading")
        manual_bl = st.file_uploader(
            "Upload draft BL",
            type=MANUAL_UPLOAD_TYPES,
            key="manual_bl_uploader",
            help="Accepted: JSON, CSV, TXT, PDF, DOCX, XLSX",
        )
        if manual_bl is not None:
            display_uploaded_attachment(manual_bl, "bl")

    current_signature = None
    if manual_si is not None and manual_bl is not None:
        current_signature = hashlib.sha256(
            manual_si.getvalue() + b"\x00" + manual_bl.getvalue()
        ).hexdigest()

    action_left, action_right, action_space = st.columns([1.4, 1, 2.6])
    compare_clicked = action_left.button(
        "Compare Uploaded Documents",
        type="primary",
        width="stretch",
        disabled=manual_si is None or manual_bl is None or not PIPELINE_AVAILABLE,
    )
    clear_clicked = action_right.button(
        "Clear Result",
        width="stretch",
        disabled=st.session_state.manual_comparison is None,
    )

    if clear_clicked:
        st.session_state.manual_comparison = None
        st.rerun()

    if compare_clicked and manual_si is not None and manual_bl is not None:
        with st.spinner("Extracting and comparing the uploaded documents..."):
            comparison = compare_uploaded_documents(
                manual_si.name,
                manual_si.getvalue(),
                manual_bl.name,
                manual_bl.getvalue(),
            )
        comparison["signature"] = current_signature
        comparison["si_filename"] = manual_si.name
        comparison["bl_filename"] = manual_bl.name
        comparison["compared_at_utc"] = datetime.now(timezone.utc).isoformat()
        st.session_state.manual_comparison = comparison

    manual_result = st.session_state.manual_comparison
    if manual_result and manual_result.get("signature") != current_signature:
        st.warning("The uploaded files changed. Select Compare Uploaded Documents to refresh the result.")
        manual_result = None

    if not PIPELINE_AVAILABLE:
        st.error(f"The document extraction pipeline is unavailable: {PIPELINE_IMPORT_ERROR}")
    elif manual_si is None or manual_bl is None:
        st.caption("Both documents are required before comparison can begin.")

    if manual_result:
        st.divider()
        st.subheader("Comparison Result")
        status_banner(manual_result)
        st.caption(
            f"SI: {manual_result['si_filename']} ({manual_result['si_format']}) · "
            f"BL: {manual_result['bl_filename']} ({manual_result['bl_format']})"
        )

        if manual_result.get("error"):
            st.error(manual_result["error"])

        if manual_result.get("rows"):
            manual_table = pd.DataFrame(manual_result["rows"])
            st.dataframe(
                manual_table[["Field", "SI Value", "BL Value", "Outcome"]],
                width="stretch",
                hide_index=True,
                column_config={
                    "Field": st.column_config.TextColumn(width="medium"),
                    "SI Value": st.column_config.TextColumn(width="large"),
                    "BL Value": st.column_config.TextColumn(width="large"),
                    "Outcome": st.column_config.TextColumn(width="medium"),
                },
            )
            evidence_details(manual_table)

            missing = manual_result.get("missing_fields", [])
            if missing:
                readable = ", ".join(field.replace("_", " ").title() for field in missing)
                st.warning(f"Values requiring human review: {readable}")

            download_result = {
                "si_filename": manual_result["si_filename"],
                "bl_filename": manual_result["bl_filename"],
                "status": manual_result["status"],
                "review_reason": manual_result.get("review_reason"),
                "has_defect": manual_result["has_defect"],
                "defect_fields": manual_result["defect_fields"],
                "missing_fields": manual_result.get("missing_fields", []),
                "compared_at_utc": manual_result["compared_at_utc"],
                "comparisons": {
                    row["Field Key"]: {
                        "si_value": row["SI Value"],
                        "bl_value": row["BL Value"],
                        "outcome": row["Outcome"],
                    }
                    for row in manual_result["rows"]
                },
            }
            result_download, csv_download = st.columns(2)
            with result_download:
                st.download_button(
                    "Download Comparison JSON",
                    data=json.dumps(download_result, indent=2, ensure_ascii=False) + "\n",
                    file_name="manual_comparison.json",
                    mime="application/json",
                    width="stretch",
                )
            with csv_download:
                st.download_button(
                    "Download Comparison CSV",
                    data=manual_table[
                        ["Field", "SI Value", "BL Value", "Outcome"]
                    ].to_csv(index=False),
                    file_name="manual_comparison.csv",
                    mime="text/csv",
                    width="stretch",
                )


# -----------------------------------------------------------------------------
# Document comparison
# -----------------------------------------------------------------------------

elif page == "Compare Documents":
    page_header(
        "Document Comparison",
        "Inspect extracted SI and draft BL values with their source evidence.",
        "Explainable verification",
    )

    comparison_ids = results_df.loc[results_df["Category"] == "BL_COMPARISON", "Email ID"].tolist()
    if not comparison_ids:
        st.info("No BL comparison requests are available. Run verification or load a result file first.")
    else:
        preferred = st.session_state.comparison_email
        default_index = comparison_ids.index(preferred) if preferred in comparison_ids else 0
        selected_id = st.selectbox("Select comparison request", comparison_ids, index=default_index)
        st.session_state.comparison_email = selected_id
        email = email_lookup[selected_id]
        result = normalise_result(selected_id, submission.get(selected_id))

        summary, issue_metric = st.columns([4, 1], vertical_alignment="center")
        summary.subheader(email.get("subject") or "No subject")
        summary.caption(f"{selected_id} · {email.get('from', 'Unknown sender')}")
        issue_metric.metric("Issues", len(result.get("defect_fields", [])))
        status_banner(result)
        st.write("")

        comparison_table, evidence = comparison_table_data(data_source, selected_id, result)
        if evidence.get("error"):
            st.warning(f"Source evidence could not be reconstructed: {evidence['error']}")
        else:
            st.caption(
                f"Source formats: SI {evidence['si_format']} · BL {evidence['bl_format']} · "
                "values compared after semantic normalization"
            )

        st.subheader("Seven-Field Comparison")
        display_table = comparison_table[["Field", "SI Value", "BL Value", "Outcome"]]
        st.dataframe(
            display_table,
            width="stretch",
            hide_index=True,
            column_config={
                "Field": st.column_config.TextColumn(width="medium"),
                "SI Value": st.column_config.TextColumn(width="large"),
                "BL Value": st.column_config.TextColumn(width="large"),
                "Outcome": st.column_config.TextColumn(width="medium"),
            },
        )
        evidence_details(comparison_table)

        action_columns = st.columns([1, 1, 3])
        if action_columns[0].button("Retry extraction", width="stretch"):
            comparison_evidence.clear()
            attachment_preview_text.clear()
            st.rerun()
        action_columns[1].button(
            "Send to review",
            width="stretch",
            on_click=navigate,
            args=("Review Queue", selected_id),
        )

        email_tab, si_tab, bl_tab = st.tabs(["Email", "Shipping Instruction", "Draft Bill of Lading"])
        with email_tab:
            st.text_area("Email body", email.get("body", ""), height=320, disabled=True)
        si_paths, bl_paths = attachment_roles(email)
        with si_tab:
            if not si_paths:
                st.warning("No Shipping Instruction attachment was found.")
            for index, path in enumerate(si_paths):
                display_attachment(inbox, path, f"si_{selected_id}_{index}")
        with bl_tab:
            if not bl_paths:
                st.warning("No draft Bill of Lading attachment was found.")
            for index, path in enumerate(bl_paths):
                display_attachment(inbox, path, f"bl_{selected_id}_{index}")


# -----------------------------------------------------------------------------
# Human review queue
# -----------------------------------------------------------------------------

elif page == "Review Queue":
    page_header(
        "Human Review Queue",
        "Confirm uncertain results, correct decisions, and preserve an audit trail.",
        "Human in the loop",
    )

    if st.session_state.review_flash:
        st.success(st.session_state.review_flash)
        st.session_state.review_flash = None

    review_df = results_df[results_df["Status"] == "NEEDS_REVIEW"].copy()
    reason_counter = Counter(review_df["Review Reason"].tolist())
    metrics = st.columns(5)
    metrics[0].metric("Awaiting Review", len(review_df))
    metrics[1].metric("Missing Attachments", reason_counter.get("missing_attachment", 0))
    metrics[2].metric("Missing Values", reason_counter.get("missing_value", 0))
    metrics[3].metric("Unreadable", reason_counter.get("unreadable", 0))
    metrics[4].metric("Wrong Document", reason_counter.get("wrong_doc_type", 0))

    reviewable_ids = review_df["Email ID"].tolist()
    requested_id = st.session_state.review_email
    if requested_id and requested_id in email_lookup and requested_id not in reviewable_ids:
        reviewable_ids = [requested_id] + reviewable_ids

    if not reviewable_ids:
        st.success("There are no cases awaiting human review.")
    else:
        st.dataframe(
            review_df[["Email ID", "Subject", "Review Reason", "Defect Fields"]],
            width="stretch",
            hide_index=True,
            height=260,
        )
        default_index = reviewable_ids.index(requested_id) if requested_id in reviewable_ids else 0
        selected_id = st.selectbox("Review a case", reviewable_ids, index=default_index)
        st.session_state.review_email = selected_id
        email = email_lookup[selected_id]
        current = normalise_result(selected_id, submission.get(selected_id))
        status_banner(current)

        context_tab, evidence_tab = st.tabs(["Email Context", "Extracted Evidence"])
        with context_tab:
            st.text_area("Email context", email.get("body", ""), height=240, disabled=True)
        with evidence_tab:
            review_table, evidence = comparison_table_data(data_source, selected_id, current)
            st.dataframe(
                review_table[["Field", "SI Value", "BL Value", "Outcome"]],
                width="stretch",
                hide_index=True,
            )
            if evidence.get("error"):
                st.warning(f"Evidence issue: {evidence['error']}")

        st.subheader("Reviewer Decision")
        with st.form(f"review_form_{selected_id}"):
            form_columns = st.columns(2)
            reviewer = form_columns[0].text_input("Reviewer name or initials")
            decision = form_columns[1].selectbox(
                "Decision", ["Keep as NEEDS_REVIEW", "Confirm OK", "Confirm MISMATCH"]
            )
            confirmed_fields = st.multiselect(
                "Confirmed mismatch fields",
                COMPARISON_FIELDS,
                default=current.get("defect_fields", []),
                format_func=lambda value: value.replace("_", " ").title(),
            )
            review_reason = st.selectbox(
                "Reason if review remains unresolved",
                sorted(VALID_REVIEW_REASONS),
                index=sorted(VALID_REVIEW_REASONS).index(current.get("review_reason"))
                if current.get("review_reason") in VALID_REVIEW_REASONS
                else 0,
                format_func=lambda value: value.replace("_", " ").title(),
            )
            note = st.text_area("Reviewer note", placeholder="Explain why the result was confirmed or corrected.")
            saved = st.form_submit_button("Save Review Decision", type="primary")

        if saved:
            if not reviewer.strip() or not note.strip():
                st.error("Reviewer name and note are required for the audit trail.")
            elif decision == "Confirm MISMATCH" and not confirmed_fields:
                st.error("Select at least one mismatch field.")
            else:
                if decision == "Confirm OK":
                    updated = {
                        "category": "BL_COMPARISON",
                        "status": "OK",
                        "review_reason": None,
                        "defect_fields": [],
                        "has_defect": False,
                    }
                elif decision == "Confirm MISMATCH":
                    updated = {
                        "category": "BL_COMPARISON",
                        "status": "MISMATCH",
                        "review_reason": None,
                        "defect_fields": confirmed_fields,
                        "has_defect": True,
                    }
                else:
                    updated = {
                        "category": "BL_COMPARISON",
                        "status": "NEEDS_REVIEW",
                        "review_reason": review_reason,
                        "defect_fields": [],
                        "has_defect": False,
                    }

                previous = submission.get(selected_id)
                st.session_state.submission[selected_id] = updated
                st.session_state.submission_source = "Reviewed in this session"
                st.session_state.review_history.append(
                    {
                        "email_id": selected_id,
                        "reviewer": reviewer.strip(),
                        "decision": updated["status"],
                        "note": note.strip(),
                        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
                        "previous_result": previous,
                        "updated_result": updated,
                    }
                )
                st.session_state.review_flash = (
                    "Review decision saved in this session. Download the revised JSON before closing the app."
                )
                st.rerun()

    if st.session_state.review_history:
        st.subheader("Review Audit Trail")
        history_rows = [
            {
                "Email ID": item["email_id"],
                "Reviewer": item["reviewer"],
                "Decision": item["decision"],
                "Note": item["note"],
                "Time (UTC)": item["timestamp_utc"],
            }
            for item in reversed(st.session_state.review_history)
        ]
        st.dataframe(pd.DataFrame(history_rows), width="stretch", hide_index=True)

    if submission:
        st.write("")
        download_col, audit_col = st.columns(2)
        with download_col:
            submission_download(submission, "review_submission_download")
        with audit_col:
            st.download_button(
                "Download Review Audit JSON",
                data=json.dumps(st.session_state.review_history, indent=2, ensure_ascii=False) + "\n",
                file_name="review_audit.json",
                mime="application/json",
                width="stretch",
            )


# -----------------------------------------------------------------------------
# Results, analytics, validation, and health
# -----------------------------------------------------------------------------

elif page == "Results & Health":
    page_header(
        "Results & Health",
        "Validate the final output, inspect operational patterns, and export results.",
        "Quality assurance",
    )

    if not submission:
        st.warning("No results are available. Run verification or upload a submission JSON.")
    else:
        summary_col, download_col = st.columns([3, 1], vertical_alignment="center")
        with summary_col:
            if validation_errors:
                st.error(f"Validation found {len(validation_errors)} issue(s).")
                with st.expander("View validation issues", expanded=True):
                    for error in validation_errors[:100]:
                        st.write(f"- {error}")
                    if len(validation_errors) > 100:
                        st.write("- Additional issues were omitted from this view.")
            else:
                st.success(f"Validation passed: all {len(submission)} email results are structurally ready.")
        with download_col:
            submission_download(submission, "results_submission_download")

        st.dataframe(
            results_df[
                ["Email ID", "Category Display", "Status Display", "Has Defect", "Defect Fields", "Review Reason"]
            ],
            width="stretch",
            hide_index=True,
            height=470,
            column_config={"Category Display": "Category", "Status Display": "Status"},
        )

        all_defect_fields: list[str] = []
        for value in submission.values():
            fields = value.get("defect_fields") or []
            if isinstance(fields, list):
                all_defect_fields.extend(fields)

        analytics_left, analytics_right = st.columns(2)
        with analytics_left:
            st.subheader("Frequently Mismatched Fields")
            if all_defect_fields:
                field_counts = pd.Series(all_defect_fields).value_counts().rename_axis("Field").reset_index(name="Count")
                field_counts["Field"] = field_counts["Field"].str.replace("_", " ").str.title()
                figure = px.bar(field_counts, x="Field", y="Count", text="Count", color_discrete_sequence=["#F35D77"])
                figure.update_layout(showlegend=False, xaxis_title="", yaxis_title="Mismatches")
                st.plotly_chart(format_plot(figure), width="stretch")
            else:
                st.info("No mismatch fields are present in the current results.")

        with analytics_right:
            st.subheader("Attachment Formats")
            extensions = [
                Path(path).suffix.lower().replace(".", "").upper() or "UNKNOWN"
                for email in emails
                for path in email.get("attachments", [])
            ]
            if extensions:
                format_counts = pd.Series(extensions).value_counts().rename_axis("Format").reset_index(name="Count")
                figure = px.bar(
                    format_counts,
                    x="Format",
                    y="Count",
                    text="Count",
                    color="Format",
                    color_discrete_sequence=["#F4A62A", "#6C89FF", "#33D69F", "#A879F7"],
                )
                figure.update_layout(showlegend=False, xaxis_title="", yaxis_title="Attachments")
                st.plotly_chart(format_plot(figure), width="stretch")

        st.subheader("System Health")
        health_data = {
            "Inbox connected": "Yes",
            "Emails loaded": len(emails),
            "Results loaded": len(submission),
            "Result source": st.session_state.submission_source,
            "Validation": "Passed" if not validation_errors else f"{len(validation_errors)} issue(s)",
            "Pipeline available": "Yes" if PIPELINE_AVAILABLE else "No",
            "Data mode": "HTTP server" if inbox.is_http else "Static bundle",
            "Application version": "2.0",
        }
        st.json(health_data)

        st.subheader("Workflow")
        st.markdown(
            """
            1. Read every inbox email.
            2. Classify it into one of five required categories.
            3. Continue only BL comparison requests to document checking.
            4. Extract the SI and draft BL fields from TXT, PDF, DOCX, or XLSX.
            5. Compare the seven shipment fields after semantic normalization.
            6. Report no mismatch, mismatch, or human review required.
            7. Validate and export one result for every email ID.
            """
        )

        if inbox.is_http:
            if st.button("Run Self-Evaluation", disabled=bool(validation_errors)):
                try:
                    with st.spinner("Submitting results for self-evaluation..."):
                        scoreboard = inbox.submit(submission)
                    st.success("Self-evaluation completed.")
                    st.json(scoreboard)
                except Exception as error:
                    st.error(f"Self-evaluation failed: {error}")
        else:
            st.info("Self-evaluation is available when the data source is the supplied HTTP server.")
